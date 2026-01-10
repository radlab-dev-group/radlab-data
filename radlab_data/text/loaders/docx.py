import docx
import langchain_core


from radlab_data.text.loaders.loader_i import LoaderI


class DOCXLoader(LoaderI):
    def __init__(self, filepath: str, options: dict = None) -> None:
        """
        :param filepath: Path to docx file
        """
        if options is None:
            options = {}

        LoaderI.__init__(self, filepath=filepath, options=options)
        self.document = docx.Document(filepath)

    def load(self) -> list:
        """
        Loads content from docx.
        :return:
        """
        if len(self.doc_pages):
            return self.doc_pages

        page_text = ""
        page_number = 0
        is_new_page = False
        proper_pages = self.options.get("merge_to_proper_pages", False)
        for p in self.document.paragraphs:
            if p.text is None or not len(p.text.strip()):
                continue

            is_new_page = self._is_new_page(p)
            if is_new_page:
                page_number += 1

            if proper_pages:
                if is_new_page:
                    lgc_d = langchain_core.documents.base.Document(
                        page_content=page_text,
                        metadata={"source": self.filepath, "page": page_number},
                    )
                    self.doc_pages.append(lgc_d)
                    page_text = p.text.strip()
                else:
                    page_text += "\n" + p.text.strip()
            else:
                page_text = p.text.strip()
                lgc_d = langchain_core.documents.base.Document(
                    page_content=page_text,
                    metadata={"source": self.filepath, "page": page_number},
                )
                self.doc_pages.append(lgc_d)
        if not is_new_page and len(page_text.strip()):
            page_number += 1
            lgc_d = langchain_core.documents.base.Document(
                page_content=page_text,
                metadata={"source": self.filepath, "page": page_number},
            )
            self.doc_pages.append(lgc_d)

        table_number = 0
        for table in self.document.tables:
            table_number += 1

            for row_number, row in enumerate(table.rows):
                for column_number, cell in enumerate(row.cells):
                    if cell.text is None or not len(cell.text.strip()):
                        continue

                    cell_text = cell.text.strip()
                    lgc_d = langchain_core.documents.base.Document(
                        page_content=cell_text,
                        metadata={
                            "source": self.filepath,
                            "table": table_number,
                            "row": row_number,
                            "column": column_number,
                        },
                    )
                    self.doc_pages.append(lgc_d)
        return self.doc_pages

    @staticmethod
    def _is_new_page(paragraph) -> bool:
        is_new_page = False
        for run in paragraph.runs:
            # Soft and hard page wrapping
            if "lastRenderedPageBreak" in run._element.xml or (
                "w:br" in run._element.xml and 'type="page"' in run._element.xml
            ):
                is_new_page = True
        return is_new_page
